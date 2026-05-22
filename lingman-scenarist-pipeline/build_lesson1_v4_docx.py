from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/maksymbabiev/Documents/СЦЕНАРИСТ")
SOURCE = ROOT / "LINGMAN_LESSON_001_TO_BE_SCRIPT_V4_NARRATIVE_REWRITE.txt"
OUT = ROOT / "LINGMAN_LESSON_001_TO_BE_V4_NARRATIVE_REWRITE.docx"


def apply_font(run, size=11, color=RGBColor(0, 0, 0), bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def is_english_line(text):
    stripped = text.strip()
    if not stripped or len(stripped) > 80:
        return False
    has_latin = any(("A" <= ch <= "Z") or ("a" <= ch <= "z") for ch in stripped)
    has_cyrillic = any("А" <= ch <= "я" or ch in "Ёё" for ch in stripped)
    return has_latin and not has_cyrillic


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

title = styles["Title"]
title.font.name = "Arial"
title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
title.font.size = Pt(26)
title.font.bold = False
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
title.paragraph_format.line_spacing = 1.15

subtitle = styles["Subtitle"]
subtitle.font.name = "Arial"
subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
subtitle.font.size = Pt(11)
subtitle.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
subtitle.paragraph_format.space_after = Pt(8)
subtitle.paragraph_format.line_spacing = 1.15

example = styles.add_style("Script Example", 1)
example.font.name = "Arial"
example._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
example._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
example._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
example.font.size = Pt(11)
example.font.bold = True
example.paragraph_format.space_before = Pt(0)
example.paragraph_format.space_after = Pt(4)
example.paragraph_format.line_spacing = 1.15

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Стр. ")
apply_font(run, size=9, color=RGBColor(0x55, 0x55, 0x55))
add_field(footer, "PAGE")
for item in footer.runs:
    apply_font(item, size=9, color=RGBColor(0x55, 0x55, 0x55))

doc.add_paragraph("Вы удаляете AM IS ARE, и английский сразу звучит сломанно", style="Title")
doc.add_paragraph("Сборка Английского по Фрейзмену | Урок 1 из 32 | V4 NARRATIVE REWRITE", style="Subtitle")

for block in SOURCE.read_text(encoding="utf-8").split("\n\n"):
    text = block.strip()
    if not text:
        continue
    paragraph = doc.add_paragraph(text, style="Script Example" if is_english_line(text) else "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.save(OUT)
print(OUT)
