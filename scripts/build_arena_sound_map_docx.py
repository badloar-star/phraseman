#!/usr/bin/env python3
"""Build the copy-friendly Arena sound map as a linear Word document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_arena_sound_map import SECTION_FILES, SoundSpec, build_sounds


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "output" / "pdf" / "zvukovaya_karta_arena_phraseman.docx"

INK = "171B21"
MUTED = "59616D"
AMBER = "B8750E"
PALE_AMBER = "F7E8C9"
SOFT = "F4F3F0"
GREEN = "DDF2E8"
GREEN_TEXT = "176544"
BLUE = "E8EFF8"
BLUE_TEXT = "365D89"

VARIANT_LABELS = {
    "A": "интерфейсный",
    "B": "игровой",
    "C": "кинематографичный",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_bottom_border(paragraph, color: str = "D9D4CC", size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(color)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in (
        ("Title", 28, INK, 0, 8),
        ("Heading 1", 17, INK, 0, 10),
        ("Heading 2", 13, INK, 6, 5),
        ("Subtitle", 11, MUTED, 0, 10),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name, size, color, bold in (
        ("Arena Eyebrow", 8, AMBER, True),
        ("Arena Meta", 8.5, MUTED, False),
        ("Arena Label", 9, AMBER, True),
        ("Arena Prompt", 9, INK, False),
        ("Arena Note", 8.5, MUTED, False),
    ):
        style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.08

    styles["Arena Eyebrow"].paragraph_format.space_after = Pt(6)
    styles["Arena Label"].paragraph_format.space_before = Pt(4)
    styles["Arena Label"].paragraph_format.space_after = Pt(2)
    styles["Arena Label"].paragraph_format.keep_with_next = True
    styles["Arena Prompt"].paragraph_format.left_indent = Cm(0.25)
    styles["Arena Prompt"].paragraph_format.right_indent = Cm(0.25)
    styles["Arena Prompt"].paragraph_format.space_before = Pt(3)
    styles["Arena Prompt"].paragraph_format.space_after = Pt(7)
    styles["Arena Prompt"].paragraph_format.keep_together = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Phraseman · Arena audio map    ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def add_running_header(section) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.text = "ЗВУКОВАЯ КАРТА АРЕНЫ"
    run = paragraph.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(AMBER)
    set_bottom_border(paragraph, "D9D4CC", "6")
    add_page_number(section.footer.paragraphs[0])


def add_cover(document: Document, sounds: list[SoundSpec]) -> None:
    paragraph = document.add_paragraph("PHRASEMAN · SOUND DESIGN", style="Arena Eyebrow")
    paragraph.paragraph_format.space_before = Pt(34)

    document.add_paragraph("Звуковая карта\nАрены", style="Title")
    document.add_paragraph(
        "Полное техническое задание на озвучку поиска соперника, столкновения, "
        "ответов, комбо, итогов, прогресса, приглашений и магазина.",
        style="Subtitle",
    )

    stats = document.add_table(rows=1, cols=4)
    stats.autofit = False
    values = (
        ("35", "звуков"),
        ("105", "промптов"),
        ("27", "триггеров подключено"),
        ("8", "нужно подключить"),
    )
    for cell, (number, label) in zip(stats.rows[0].cells, values):
        set_cell_shading(cell, SOFT)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(number + "\n")
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(17)
        r.font.color.rgb = RGBColor.from_string(AMBER)
        r = p.add_run(label)
        r.font.name = "Arial"
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(20)
    note.paragraph_format.space_after = Pt(7)
    run = note.add_run("Версия для нормального копирования")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(GREEN_TEXT)
    set_paragraph_shading(note, GREEN)

    document.add_paragraph(
        "Каждый вариант A / B / C записан отдельным обычным абзацем. "
        "Метки вариантов больше не стоят сбоку и не вклиниваются в текст при выделении. "
        "Для работы с промптами используйте DOCX; PDF оставлен как удобная версия для просмотра.",
        style="Arena Note",
    )
    document.add_paragraph("Обновлено: 25 августа 2026", style="Arena Note")


def add_usage(document: Document) -> None:
    document.add_page_break()
    document.add_paragraph("Как пользоваться", style="Heading 1")
    instructions = (
        "Откройте DOCX, найдите нужный mp3 и скопируйте весь серый абзац под вариантом A, B или C.",
        "A — сухой интерфейсный звук; B — более игровой; C — кинематографичный.",
        "Длительность, громкость, cooldown и приоритет уже согласованы с каталогом Арены.",
        "Статус «нужно добавить триггер» означает, что звук описан, но вызов в интерфейсе ещё не подключён.",
    )
    for index, text in enumerate(instructions, 1):
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(f"{index}. ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(AMBER)
        p.add_run(text)

    document.add_paragraph("Разделы", style="Heading 2")
    for section, files in SECTION_FILES:
        p = document.add_paragraph(style="Arena Note")
        p.add_run(f"{section}  ").bold = True
        p.add_run(f"{len(files)} зв.")


def add_sound(document: Document, index: int, sound: SoundSpec) -> None:
    heading = document.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    number = heading.add_run(f"{index:02d}  ")
    number.font.color.rgb = RGBColor.from_string(AMBER)
    heading.add_run(sound.file)

    status = "триггер подключён" if sound.integration == "wired" else "нужно добавить триггер"
    meta = document.add_paragraph(style="Arena Meta")
    meta.paragraph_format.keep_with_next = True
    meta.add_run(
        f"{sound.duration_ms / 1000:g} с длина   ·   {sound.volume:g} громкость   ·   "
        f"{sound.cooldown_ms} мс cooldown   ·   {sound.priority} приоритет   ·   {status}   ·   файл ожидается"
    )
    set_paragraph_shading(meta, GREEN if sound.integration == "wired" else PALE_AMBER)

    moment = document.add_paragraph()
    moment.paragraph_format.keep_with_next = True
    r = moment.add_run("Момент: ")
    r.bold = True
    moment.add_run(sound.moment)

    trigger = document.add_paragraph()
    trigger.paragraph_format.keep_with_next = True
    r = trigger.add_run("Точка: ")
    r.bold = True
    trigger.add_run(sound.trigger)

    for variant in ("A", "B", "C"):
        label = document.add_paragraph(
            f"{variant} — {VARIANT_LABELS[variant]}", style="Arena Label"
        )
        label.paragraph_format.keep_with_next = True
        prompt = document.add_paragraph(sound.prompts[variant], style="Arena Prompt")
        set_paragraph_shading(prompt, SOFT)

    divider = document.add_paragraph()
    divider.paragraph_format.space_after = Pt(7)
    set_bottom_border(divider)


def add_summary(document: Document, sounds: list[SoundSpec]) -> None:
    document.add_page_break()
    document.add_paragraph("Сводная карта файлов", style="Heading 1")
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ("Файл", "Длина", "Громк.", "Cooldown", "Статус")
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, color="FFFFFF")
        set_cell_shading(cell, INK)
    set_repeat_table_header(table.rows[0])

    for sound in sounds:
        row = table.add_row()
        status = "подключён" if sound.integration == "wired" else "добавить триггер"
        values = (
            sound.file,
            f"{sound.duration_ms / 1000:g} с",
            f"{sound.volume:g}",
            f"{sound.cooldown_ms} мс",
            status,
        )
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)
        set_cell_shading(row.cells[-1], GREEN if sound.integration == "wired" else PALE_AMBER)


def build_document(sounds: list[SoundSpec]) -> Document:
    document = Document()
    configure_document(document)
    add_running_header(document.sections[0])
    document.core_properties.title = "Звуковая карта Арены Phraseman"
    document.core_properties.subject = "35 звуков и 105 промптов для Arena"
    document.core_properties.author = "Phraseman"
    document.core_properties.keywords = "Phraseman, Arena, sound design, audio map"

    add_cover(document, sounds)
    add_usage(document)

    by_file = {sound.file: sound for sound in sounds}
    index = 0
    for section_index, (section_title, files) in enumerate(SECTION_FILES):
        if section_index == 0:
            document.add_page_break()
        document.add_paragraph(section_title, style="Heading 1")
        for file in files:
            index += 1
            add_sound(document, index, by_file[file])

    add_summary(document, sounds)
    return document


def main() -> int:
    sounds = build_sounds()
    document = build_document(sounds)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(f"PASS: DOCX created - {OUTPUT_PATH}")
    print(f"PASS: {len(sounds)} sounds, {sum(len(s.prompts) for s in sounds)} prompts")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
