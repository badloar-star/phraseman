from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/maksymbabiev/Documents/СЦЕНАРИСТ")
SOURCE = ROOT / "LINGMAN_LESSON_001_TO_BE_SCRIPT_V3_FULL_REWRITE.txt"
OUT = ROOT / "LINGMAN_LESSON_001_TO_BE_V3_FULL_REWRITE.docx"


def set_font(run, *, name="Arial", size=11, color=RGBColor(0, 0, 0), bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def set_paragraph_spacing(paragraph, *, before=0, after=8, line=1.15):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_field(paragraph, field):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


def is_english_example(text):
    stripped = text.strip()
    if not stripped or len(stripped) > 90:
        return False
    latin_letters = sum(1 for ch in stripped if ("A" <= ch <= "Z") or ("a" <= ch <= "z"))
    cyrillic_letters = sum(1 for ch in stripped if "А" <= ch <= "я" or ch == "ё" or ch == "Ё")
    if latin_letters == 0 or cyrillic_letters > 0:
        return False
    starters = (
        "I ", "I'm", "You ", "You're", "He ", "He's", "She ", "She's",
        "It ", "It's", "We ", "We're", "They ", "They're", "Masha ",
        "My ", "The ", "This ", "Are ", "Do ",
    )
    return stripped.startswith(starters) or stripped in {"Am", "Is", "Are", "I", "He", "She", "It", "You", "We", "They"}


doc = Document()

section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

title = styles["Title"]
title.font.name = "Arial"
title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
title.font.size = Pt(26)
title.font.color.rgb = RGBColor(0, 0, 0)
title.font.bold = False
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
title.paragraph_format.line_spacing = 1.15

subtitle = styles["Subtitle"]
subtitle.font.name = "Arial"
subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
subtitle.font.size = Pt(11)
subtitle.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
subtitle.paragraph_format.space_before = Pt(0)
subtitle.paragraph_format.space_after = Pt(8)
subtitle.paragraph_format.line_spacing = 1.15

example = styles.add_style("Script Example", 1)
example.font.name = "Arial"
example._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
example._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
example._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
example.font.size = Pt(11)
example.font.color.rgb = RGBColor(0, 0, 0)
example.font.bold = True
example.paragraph_format.space_before = Pt(0)
example.paragraph_format.space_after = Pt(4)
example.paragraph_format.line_spacing = 1.15

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_run = footer.add_run("Стр. ")
set_font(footer_run, size=9, color=RGBColor(0x55, 0x55, 0x55))
add_field(footer, "PAGE")
for run in footer.runs:
    set_font(run, size=9, color=RGBColor(0x55, 0x55, 0x55))

doc.add_paragraph("Вы удаляете AM IS ARE, и английский сразу звучит сломанно", style="Title")
doc.add_paragraph(
    "Сборка Английского по Фрейзмену | Урок 1 из 32 | V3 FULL REWRITE",
    style="Subtitle",
)

for block in SOURCE.read_text(encoding="utf-8").split("\n\n"):
    text = block.strip()
    if not text:
        continue
    paragraph = doc.add_paragraph(text, style="Script Example" if is_english_example(text) else "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, after=4 if is_english_example(text) else 8)

doc.save(OUT)
print(OUT)
