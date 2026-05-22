from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/maksymbabiev/Documents/СЦЕНАРИСТ")
SOURCE = ROOT / "LINGMAN_SCRIPT_002_7_RUSSIAN_ENGLISH_MISTAKES_10_MIN.txt"
OUT = ROOT / "LINGMAN_SCRIPT_002_7_RUSSIAN_ENGLISH_MISTAKES_10_MIN.docx"


def set_spacing(style, *, after_pt=6, before_pt=0, line=1.10):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before_pt)
    fmt.space_after = Pt(after_pt)
    fmt.line_spacing = line


def add_field(paragraph, field):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


def is_example(paragraph_text):
    stripped = paragraph_text.strip()
    if not stripped:
        return False
    if stripped.startswith("“") and stripped.endswith("”") and any(ch.isascii() and ch.isalpha() for ch in stripped):
        return True
    english_markers = (
        "I ", "Do ", "Are ", "Can ", "Call ", "She ", "He ", "This ",
        "The ", "Listen ", "Talk ", "Speak ", "Write ",
    )
    return stripped.startswith(english_markers) and any(ch.isascii() and ch.isalpha() for ch in stripped)


doc = Document()

section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
set_spacing(normal, after_pt=6, before_pt=0, line=1.10)

title = styles["Title"]
title.font.name = "Calibri"
title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
title.font.size = Pt(22)
title.font.bold = True
title.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
set_spacing(title, after_pt=4, before_pt=0, line=1.10)

subtitle = styles["Subtitle"]
subtitle.font.name = "Calibri"
subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
subtitle.font.size = Pt(11)
subtitle.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
set_spacing(subtitle, after_pt=14, before_pt=0, line=1.10)

example = styles.add_style("Script Example", 1)
example.font.name = "Consolas"
example._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
example.font.size = Pt(11)
example.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
example.paragraph_format.left_indent = Inches(0.25)
example.paragraph_format.space_before = Pt(0)
example.paragraph_format.space_after = Pt(6)
example.paragraph_format.line_spacing = 1.10

header = section.header.paragraphs[0]
header.text = "Professor Lingman - сценарий"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Стр. ")
add_field(footer, "PAGE")
for run in footer.runs:
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_paragraph("7 ошибок в английском, которые сразу выдают русскоязычного", style="Title")
doc.add_paragraph("Телесуфлерный сценарий на 10 минут | Professor Lingman", style="Subtitle")

for block in SOURCE.read_text(encoding="utf-8").split("\n\n"):
    text = block.strip()
    if not text:
        continue
    style = "Script Example" if is_example(text) else "Normal"
    para = doc.add_paragraph(text, style=style)
    if style == "Script Example":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.save(OUT)
print(OUT)
